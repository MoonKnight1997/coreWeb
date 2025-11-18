"""
Document Processing Service

Handles document upload, validation, and text extraction for various file formats.
"""

import os
import tempfile
import uuid
from typing import Tuple, Optional
from werkzeug.utils import secure_filename
from werkzeug.datastructures import FileStorage

from app.config import current_config
from app.utils.logger import logger


class DocumentProcessingError(Exception):
    """Raised when document processing fails."""
    pass


class DocumentProcessor:
    """Handles document processing operations."""

    @staticmethod
    def save_upload(file: FileStorage) -> Tuple[str, str]:
        """
        Save uploaded file securely to temporary location.

        Args:
            file: Uploaded file from Flask request

        Returns:
            Tuple of (file_path, original_filename)

        Raises:
            DocumentProcessingError: If file save fails
        """
        try:
            # Generate unique filename
            filename = secure_filename(file.filename)
            unique_id = str(uuid.uuid4())[:8]
            temp_filename = f"{unique_id}_{filename}"

            # Save to temporary location
            temp_path = os.path.join(tempfile.gettempdir(), temp_filename)
            file.save(temp_path)

            logger.info(
                'File uploaded successfully',
                filename=filename,
                temp_path=temp_path,
                file_size=os.path.getsize(temp_path)
            )

            return temp_path, filename

        except Exception as e:
            logger.error('File upload failed', error=str(e), filename=file.filename)
            raise DocumentProcessingError(f'Failed to save file: {str(e)}')

    @staticmethod
    def extract_text(filepath: str) -> str:
        """
        Extract text from document file.

        Supports: TXT, PDF, DOCX

        Args:
            filepath: Path to file

        Returns:
            Extracted text content

        Raises:
            DocumentProcessingError: If extraction fails
        """
        ext = filepath.rsplit('.', 1)[1].lower() if '.' in filepath else ''

        try:
            if ext == 'txt':
                return DocumentProcessor._extract_from_txt(filepath)
            elif ext == 'pdf':
                return DocumentProcessor._extract_from_pdf(filepath)
            elif ext in ['docx', 'doc']:
                return DocumentProcessor._extract_from_docx(filepath)
            else:
                raise DocumentProcessingError(f'Unsupported file type: .{ext}')

        except DocumentProcessingError:
            raise
        except Exception as e:
            logger.error('Text extraction failed', filepath=filepath, error=str(e))
            raise DocumentProcessingError(f'Failed to extract text: {str(e)}')

    @staticmethod
    def _extract_from_txt(filepath: str) -> str:
        """Extract text from TXT file."""
        try:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()

            if not text.strip():
                raise DocumentProcessingError('Text file is empty')

            logger.debug('Extracted text from TXT', filepath=filepath, length=len(text))
            return text

        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f'TXT extraction failed: {str(e)}')

    @staticmethod
    def _extract_from_pdf(filepath: str) -> str:
        """Extract text from PDF file."""
        try:
            import PyPDF2

            text = ''
            with open(filepath, 'rb') as f:
                pdf_reader = PyPDF2.PdfReader(f)

                if len(pdf_reader.pages) == 0:
                    raise DocumentProcessingError('PDF has no pages')

                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + '\n'
                    except Exception as e:
                        logger.warning(
                            'Failed to extract page from PDF',
                            page_num=page_num,
                            error=str(e)
                        )
                        continue

            if not text.strip():
                raise DocumentProcessingError(
                    'PDF text extraction failed. The PDF may be scanned images '
                    'or encrypted. Please use a text-based PDF or extract the text manually.'
                )

            logger.debug('Extracted text from PDF', filepath=filepath, length=len(text))
            return text

        except DocumentProcessingError:
            raise
        except ImportError:
            raise DocumentProcessingError('PDF processing library not available')
        except Exception as e:
            raise DocumentProcessingError(f'PDF extraction failed: {str(e)}')

    @staticmethod
    def _extract_from_docx(filepath: str) -> str:
        """Extract text from DOCX file."""
        try:
            import docx

            doc = docx.Document(filepath)

            # Extract text from paragraphs
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)

            text = '\n'.join(paragraphs)

            if not text.strip():
                raise DocumentProcessingError('DOCX file appears to be empty')

            logger.debug('Extracted text from DOCX', filepath=filepath, length=len(text))
            return text

        except DocumentProcessingError:
            raise
        except ImportError:
            raise DocumentProcessingError('DOCX processing library not available')
        except Exception as e:
            raise DocumentProcessingError(f'DOCX extraction failed: {str(e)}')

    @staticmethod
    def cleanup(filepath: str) -> None:
        """
        Clean up temporary file.

        Args:
            filepath: Path to file to delete
        """
        try:
            if os.path.exists(filepath):
                os.remove(filepath)
                logger.debug('Cleaned up temporary file', filepath=filepath)
        except Exception as e:
            logger.warning('Failed to cleanup file', filepath=filepath, error=str(e))

    @staticmethod
    def validate_text_length(text: str, min_length: int = 100) -> None:
        """
        Validate that text meets minimum length requirement.

        Args:
            text: Text to validate
            min_length: Minimum required length

        Raises:
            DocumentProcessingError: If text is too short
        """
        if len(text.strip()) < min_length:
            raise DocumentProcessingError(
                f'Document text is too short. Minimum {min_length} characters required, '
                f'got {len(text.strip())}.'
            )

    @staticmethod
    def process_upload(file: Optional[FileStorage], text: Optional[str]) -> str:
        """
        Process either file upload or pasted text.

        Args:
            file: Uploaded file (optional)
            text: Pasted text (optional)

        Returns:
            Extracted document text

        Raises:
            DocumentProcessingError: If processing fails
        """
        temp_filepath = None

        try:
            # Handle file upload
            if file and file.filename:
                temp_filepath, original_filename = DocumentProcessor.save_upload(file)
                document_text = DocumentProcessor.extract_text(temp_filepath)

                logger.info(
                    'Document processed from file',
                    filename=original_filename,
                    text_length=len(document_text)
                )

                return document_text

            # Handle pasted text
            elif text and text.strip():
                DocumentProcessor.validate_text_length(text)

                logger.info(
                    'Document processed from pasted text',
                    text_length=len(text)
                )

                return text.strip()

            else:
                raise DocumentProcessingError(
                    'No document provided. Please upload a file or paste text.'
                )

        finally:
            # Always cleanup temporary file
            if temp_filepath:
                DocumentProcessor.cleanup(temp_filepath)


# Convenience function
def process_document(file: Optional[FileStorage], text: Optional[str]) -> str:
    """
    Process document from file or text.

    Args:
        file: Uploaded file
        text: Pasted text

    Returns:
        Extracted document text

    Raises:
        DocumentProcessingError: If processing fails
    """
    return DocumentProcessor.process_upload(file, text)
